# -*- coding: utf-8 -*-
import datetime
from urllib.parse import quote

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.table import Table, TableStyleInfo
from rest_framework.decorators import action
from rest_framework.request import Request

from coreadmin.utils.import_export import import_to_data
from coreadmin.utils.json_response import DetailResponse, SuccessResponse, ErrorResponse
from coreadmin.utils.request_util import get_verbose_name
from coreadmin.system.tasks import async_export_data
from coreadmin.system.models import DownloadCenter


class ImportSerializerMixin:
    """
    自定义导入模板、导入功能
    """

    # 导入字段
    import_field_dict = {}
    # 导入序列化器
    import_serializer_class = None
    # 表格表头最大宽度，默认50个字符
    export_column_width = 50

    def is_number(self,num):
        try:
            float(num)
            return True
        except ValueError:
            pass

        try:
            import unicodedata
            unicodedata.numeric(num)
            return True
        except (TypeError, ValueError):
            pass
        return False

    def get_string_len(self, string):
        """
        获取字符串最大长度
        :param string:
        :return:
        """
        length = 4
        if string is None:
            return length
        if self.is_number(string):
            return length
        for char in string:
            length += 2.1 if ord(char) > 256 else 1
        return round(length, 1) if length <= self.export_column_width else self.export_column_width

    @action(methods=['get','post'],detail=False)
    @transaction.atomic  # Django 事务,防止出错
    def import_data(self, request: Request, *args, **kwargs):
        """
        导入模板
        :param request:
        :param args:
        :param kwargs:
        :return:
        """
        # DRF also maps HEAD to this action; only GET may build a template.
        if request.method != 'GET':
            return ErrorResponse(msg='Path-based import is temporarily disabled.', status=405)
        assert self.import_field_dict, "'%s' 请配置对应的导出模板字段。" % self.__class__.__name__
        # 导出模板
        if request.method == "GET":
            # 示例数据
            queryset = self.filter_queryset(self.get_queryset())
            # 导出excel 表
            response = HttpResponse(content_type="application/msexcel")
            response["Access-Control-Expose-Headers"] = f"Content-Disposition"
            response[
                "Content-Disposition"
            ] = f'attachment;filename={quote(str(f"导入{get_verbose_name(queryset)}模板.xlsx"))}'
            wb = Workbook()
            ws1 = wb.create_sheet("data", 1)
            ws1.sheet_state = "hidden"
            ws = wb.active
            row = get_column_letter(len(self.import_field_dict) + 1)
            column = 10
            header_data = [
                "序号",
            ]
            validation_data_dict = {}
            for index, ele in enumerate(self.import_field_dict.values()):
                if isinstance(ele, dict):
                    header_data.append(ele.get("title"))
                    choices = ele.get("choices", {})
                    if choices.get("data"):
                        data_list = []
                        data_list.extend(choices.get("data").keys())
                        validation_data_dict[ele.get("title")] = data_list
                    elif choices.get("queryset") and choices.get("values_name"):
                        data_list = choices.get("queryset").values_list(choices.get("values_name"), flat=True)
                        validation_data_dict[ele.get("title")] = list(data_list)
                    else:
                        continue
                    column_letter = get_column_letter(len(validation_data_dict))
                    dv = DataValidation(
                        type="list",
                        formula1=f"{quote_sheetname('data')}!${column_letter}$2:${column_letter}${len(validation_data_dict[ele.get('title')]) + 1}",
                        allow_blank=True,
                    )
                    ws.add_data_validation(dv)
                    dv.add(f"{get_column_letter(index + 2)}2:{get_column_letter(index + 2)}1048576")
                else:
                    header_data.append(ele)
            # 添加数据列
            ws1.append(list(validation_data_dict.keys()))
            for index, validation_data in enumerate(validation_data_dict.values()):
                for inx, ele in enumerate(validation_data):
                    ws1[f"{get_column_letter(index + 1)}{inx + 2}"] = ele
            # 插入导出模板正式数据
            df_len_max = [self.get_string_len(ele) for ele in header_data]
            ws.append(header_data)
            # 　更新列宽
            for index, width in enumerate(df_len_max):
                ws.column_dimensions[get_column_letter(index + 1)].width = width
            tab = Table(displayName="Table1", ref=f"A1:{row}{column}")  # 名称管理器
            style = TableStyleInfo(
                name="TableStyleLight11",
                showFirstColumn=True,
                showLastColumn=True,
                showRowStripes=True,
                showColumnStripes=True,
            )
            tab.tableStyleInfo = style
            ws.add_table(tab)
            wb.save(response)
            return response
        else:
            # 从excel中组织对应的数据结构，然后使用序列化器保存
            queryset = self.filter_queryset(self.get_queryset())
            # 获取多对多字段
            m2m_fields = [
                ele.name
                for ele in queryset.model._meta.get_fields()
                if hasattr(ele, "many_to_many") and ele.many_to_many == True
            ]
            import_field_dict = {'id':'更新主键(勿改)',**self.import_field_dict}
            data = import_to_data(request.data.get("url"), import_field_dict, m2m_fields)
            for ele in data:
                filter_dic = {'id':ele.get('id')}
                instance = filter_dic and queryset.filter(**filter_dic).first()
                serializer = self.import_serializer_class(instance, data=ele, request=request)
                serializer.is_valid(raise_exception=True)
                serializer.save()
            return DetailResponse(msg=f"导入成功！")

    @action(methods=['get'],detail=False)
    def update_template(self,request):
        queryset, objects, data, columns = self.workbook_projection(
            request, self.import_serializer_class, self.import_field_dict)
        workbook = Workbook()
        sheet = workbook.active
        validation_sheet = workbook.create_sheet('data')
        validation_sheet.sheet_state = 'hidden'
        sheet.append(['序号', '更新主键(勿改)', *[title for _, title, _ in columns]])
        for number, (obj, item) in enumerate(zip(objects, data), 1):
            # Structural identity comes only from the authorized row, even if
            # RESTQL omitted id from its presentation serializer.
            sheet.append([number, obj.pk, *[self.workbook_cell(item.get(key)) for key, _, _ in columns]])
        validation_column = 0
        for business_column, (key, title, specification) in enumerate(columns, 3):
            choices = specification.get('choices', {}) if isinstance(specification, dict) else {}
            values = self.workbook_choice_values(key, choices)
            if not values:
                continue
            validation_column += 1
            letter = get_column_letter(validation_column)
            validation_sheet.cell(1, validation_column, title)
            for row, value in enumerate(values, 2):
                validation_sheet.cell(row, validation_column, self.workbook_cell(value))
            validation = DataValidation(type='list', allow_blank=True,
                formula1=f"{quote_sheetname('data')}!${letter}$2:${letter}${len(values) + 1}")
            sheet.add_data_validation(validation)
            target = get_column_letter(business_column)
            validation.add(f'{target}2:{target}1048576')
        return self.finish_workbook(workbook, queryset)


class ExportSerializerMixin:
    """
    自定义导出功能
    """

    # 导出字段
    export_field_label = []
    # 导出序列化器
    export_serializer_class = None
    # 表格表头最大宽度，默认50个字符
    export_column_width = 50

    def is_number(self,num):
        try:
            float(num)
            return True
        except ValueError:
            pass

        try:
            import unicodedata
            unicodedata.numeric(num)
            return True
        except (TypeError, ValueError):
            pass
        return False

    def get_string_len(self, string):
        """
        获取字符串最大长度
        :param string:
        :return:
        """
        length = 4
        if string is None:
            return length
        if self.is_number(string):
            return length
        for char in string:
            length += 2.1 if ord(char) > 256 else 1
        return round(length, 1) if length <= self.export_column_width else self.export_column_width

    @action(methods=['get'],detail=False)
    def export_data(self, request: Request, *args, **kwargs):
        """
        导出功能
        :param request:
        :param args:
        :param kwargs:
        :return:
        """
        queryset, _, data, columns = self.workbook_projection(
            request, self.export_serializer_class, self.export_field_label)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['序号', *[title for _, title, _ in columns]])
        for number, item in enumerate(data, 1):
            sheet.append([number, *[self.workbook_cell(item.get(key)) for key, _, _ in columns]])
        return self.finish_workbook(workbook, queryset)

    def workbook_projection(self, request, serializer_class, configured_columns):
        """Shared authorization projection; the two file layouts remain distinct."""
        from rest_framework.exceptions import MethodNotAllowed
        if not serializer_class or not configured_columns:
            raise MethodNotAllowed(request.method)
        queryset = self.filter_queryset(self.get_queryset())
        objects = list(queryset)
        serializer = serializer_class(objects, many=True, request=request)
        readable = self.field_policy.query_fields()
        columns = []
        for key, specification in configured_columns.items():
            output = specification.get('display', key) if isinstance(specification, dict) else key
            if output != 'id' and output in readable and output in serializer.child.fields:
                title = specification.get('title', key) if isinstance(specification, dict) else specification
                columns.append((output, title, specification))
        return queryset, objects, serializer.data, columns

    def workbook_choice_values(self, output, choices):
        if choices.get('data'):
            return list(choices['data'])
        queryset = choices.get('queryset')
        field = choices.get('values_name')
        if queryset is None or not field:
            return []
        # The current queryset-backed template choices are User.dept/role.
        # Reuse their registered child policies; configuration is not a grant.
        from coreadmin.access.fields import RELATIONS, FieldPolicy
        from coreadmin.access.context import AccessContext
        from coreadmin.access.registry import REGISTRY
        from coreadmin.system.models import Dept, Role
        binding = RELATIONS.get(self.field_policy.resource, {}).get(output)
        models = {'dept': Dept, 'role': Role}
        if not binding or models.get(binding[0]) is not queryset.model:
            return []
        context = AccessContext(self.request.user, REGISTRY[binding[0], 'retrieve', 'GET'])
        if not context.allowed():
            return []
        policy = FieldPolicy(context, queryset.model)
        if field not in policy.ceiling('read') or '__' in field:
            return []
        return [getattr(obj, field) for obj in context.scope(queryset)
                if field in policy.allowed(obj)]

    @staticmethod
    def workbook_cell(value):
        return str(value) if isinstance(value, (list, dict, tuple)) else value

    def finish_workbook(self, workbook, queryset):
        """Common formatting only: preserve business filename, widths and table."""
        sheet = workbook.active
        for index, cells in enumerate(sheet.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = max(
                self.get_string_len('' if cell.value is None else str(cell.value)) for cell in cells)
        table = Table(displayName='Table',
                      ref=f'A1:{get_column_letter(sheet.max_column)}{max(2, sheet.max_row)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight11', showFirstColumn=True,
            showLastColumn=True, showRowStripes=True, showColumnStripes=True)
        sheet.add_table(table)
        response = HttpResponse(content_type='application/msexcel')
        response['Content-Disposition'] = f'attachment;filename={quote(f"导出{get_verbose_name(queryset)}.xlsx")}'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'
        workbook.save(response)
        return response
