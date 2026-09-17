"""Workbook compatibility must not broaden canonical row/field permissions."""
from io import BytesIO
from unittest.mock import patch
from urllib.parse import unquote

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from coreadmin.foundation_tests.access_fixtures import grant
from coreadmin.system.models import Dept, Users
from coreadmin.system.views.user import UserViewSet


class WorkbookContractTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.actor = Users.objects.create(username='workbook-actor')
        cls.admin = Users.objects.create(username='workbook-admin', is_superuser=True)
        cls.own = Users.objects.create(username='own-workbook', name='Visible name',
            email='HIDDEN_EMAIL@example.test', creator=cls.actor)
        cls.second = Users.objects.create(username='second-workbook', creator=cls.actor)
        cls.other = Users.objects.create(username='OUT_OF_SCOPE_USER')

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.actor)

    def workbook(self, route, method='get', query=None):
        response = getattr(self.client, method)(route, query or {})
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return response, load_workbook(BytesIO(response.content))

    def allow(self, action, fields):
        grant(self.actor, 'user:' + action, 'Users', fields=fields, scope=0)

    def assert_style(self, sheet, columns, rows):
        table = sheet.tables['Table']
        self.assertEqual(table.tableStyleInfo.name, 'TableStyleLight11')
        self.assertEqual(table.ref, f'A1:{columns}{rows}')
        for column in range(1, sheet.max_column + 1):
            from openpyxl.utils import get_column_letter
            self.assertGreater(sheet.column_dimensions[get_column_letter(column)].width, 0)

    def test_export_headers_sequence_scope_filename_and_style(self):
        self.allow('Export', ['username', 'name', 'dept_owner'])
        response, workbook = self.workbook('/api/system/user/export_data/')
        rows = list(workbook.active.values)
        self.assertEqual(rows[0], ('序号', '用户账号', '用户名称'))
        self.assertEqual([row[0] for row in rows[1:]], [1, 2])
        self.assertEqual({row[1] for row in rows[1:]}, {self.own.username, self.second.username})
        self.assertNotIn(self.other.username, str(rows))
        self.assertNotIn(self.own.email, str(rows))
        self.assertNotIn('部门负责人', str(rows))
        self.assertEqual(unquote(response['Content-Disposition']),
            f'attachment;filename=导出{Users._meta.verbose_name}.xlsx')
        self.assertEqual(response['Access-Control-Expose-Headers'], 'Content-Disposition')
        self.assertEqual(workbook.sheetnames, ['Sheet'])
        self.assert_style(workbook.active, 'C', 3)

    def test_export_get_post_have_identical_schema_rows_and_filename(self):
        self.allow('Export', ['username'])
        get_response, get_book = self.workbook('/api/system/user/export_data/')
        post_response, post_book = self.workbook('/api/system/user/export/', 'post')
        self.assertEqual(list(get_book.active.values), list(post_book.active.values))
        self.assertEqual(get_response['Content-Disposition'], post_response['Content-Disposition'])
        self.assertEqual(get_book.active.tables['Table'].ref, post_book.active.tables['Table'].ref)

    def test_export_action_deny_and_query_policy_still_apply(self):
        self.assertEqual(self.client.get('/api/system/user/export_data/').status_code, 403)
        self.assertEqual(self.client.post('/api/system/user/export/').status_code, 403)
        self.allow('Export', ['username'])
        for query in ({'email': self.own.email}, {'ordering': 'email'}, {'query': '{email}'}):
            self.assertEqual(self.client.get('/api/system/user/export_data/', query).status_code, 400)

    def test_template_structural_identity_scope_and_business_contract(self):
        self.allow('UpdateTemplate', ['username'])
        response, workbook = self.workbook('/api/system/user/update_template/', query={'query': '{username}'})
        rows = list(workbook.active.values)
        self.assertEqual(rows[0], ('序号', '更新主键(勿改)', '登录账号'))
        self.assertEqual([row[0] for row in rows[1:]], [1, 2])
        self.assertEqual({row[1]: row[2] for row in rows[1:]},
            {self.own.pk: self.own.username, self.second.pk: self.second.username})
        self.assertNotIn(self.other.pk, [row[1] for row in rows[1:]])
        self.assertNotIn(self.own.email, str(rows))
        self.assertEqual(workbook['data'].sheet_state, 'hidden')
        self.assertEqual(unquote(response['Content-Disposition']),
            f'attachment;filename=导出{Users._meta.verbose_name}.xlsx')
        self.assert_style(workbook.active, 'C', 3)

    def test_template_static_choices_validation_range_and_hidden_data(self):
        self.allow('UpdateTemplate', ['username', 'gender'])
        _, workbook = self.workbook('/api/system/user/update_template/')
        self.assertEqual(list(workbook.active.values)[0], ('序号', '更新主键(勿改)', '登录账号', '用户性别'))
        self.assertEqual(workbook['data'].sheet_state, 'hidden')
        self.assertEqual(list(workbook['data'].values), [('用户性别',), ('未知',), ('男',), ('女',)])
        validations = workbook.active.data_validations.dataValidation
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].type, 'list')
        self.assertEqual(validations[0].formula1, "'data'!$A$2:$A$4")
        self.assertEqual(str(validations[0].sqref), 'D2:D1048576')
        self.assert_style(workbook.active, 'D', 3)

    def test_unreadable_choices_not_evaluated_or_exposed(self):
        self.allow('UpdateTemplate', ['username'])
        configuration = {
            'username': '登录账号',
            'email': {'title': 'Hidden', 'choices': {'data': {'SECRET_CHOICE': 1}}},
            'dept': {'title': 'Hidden dept', 'choices': {
                'queryset': Dept.objects.all(), 'values_name': 'name'}},
        }
        with patch.object(UserViewSet, 'import_field_dict', configuration), patch.object(
                UserViewSet, 'workbook_choice_values', wraps=UserViewSet.workbook_choice_values) as choices:
            # Only the visible scalar column reaches choice resolution.
            choices.side_effect = lambda key, spec: [] if key == 'username' else self.fail('Hidden choices evaluated')
            _, workbook = self.workbook('/api/system/user/update_template/')
        self.assertEqual(choices.call_count, 1)
        self.assertNotIn('SECRET_CHOICE', str([list(sheet.values) for sheet in workbook]))
        self.assertEqual(len(workbook.active.data_validations.dataValidation), 0)

    def test_queryset_choices_use_target_scope_and_target_fields(self):
        visible = Dept.objects.create(name='AUTHORIZED_CHOICE')
        hidden = Dept.objects.create(name='HIDDEN_CHOICE')
        self.allow('UpdateTemplate', ['username', 'dept'])
        grant(self.actor, 'dept:Retrieve', 'Dept', fields=['name'], scope=4, departments=[visible])
        grant(self.actor, 'dept:Retrieve', 'Dept', fields=['key'], scope=4, departments=[hidden])
        _, workbook = self.workbook('/api/system/user/update_template/')
        values = str(list(workbook['data'].values))
        self.assertIn(visible.name, values)
        self.assertNotIn(hidden.name, values)
        self.assertEqual(workbook.active.data_validations.dataValidation[0].formula1, "'data'!$A$2:$A$2")

    def test_queryset_choices_without_related_action_are_empty(self):
        Dept.objects.create(name='NO_ACTION_CHOICE')
        self.allow('UpdateTemplate', ['username', 'dept'])
        _, workbook = self.workbook('/api/system/user/update_template/')
        self.assertNotIn('NO_ACTION_CHOICE', str(list(workbook['data'].values)))
        self.assertFalse(workbook.active.data_validations.dataValidation)

    def test_dept_update_template_retains_existing_contract(self):
        own = Dept.objects.create(name='Own dept', key='own-key', creator=self.actor)
        other = Dept.objects.create(name='Other dept')
        grant(self.actor, 'dept:UpdateTemplate', 'Dept', fields=['name', 'key'], scope=0)
        _, workbook = self.workbook('/api/system/dept/update_template/')
        rows = list(workbook.active.values)
        self.assertEqual(rows[0][:2], ('序号', '更新主键(勿改)'))
        self.assertEqual(rows[1][0:2], (1, own.pk))
        self.assertEqual(len(rows), 2)
        self.assertNotIn(other.name, str(rows))
        self.assertEqual(workbook['data'].sheet_state, 'hidden')
        self.assert_style(workbook.active, 'D', 2)

    def test_missing_configuration_does_not_enable_workbook(self):
        self.client.force_authenticate(self.admin)
        with patch.object(UserViewSet, 'export_serializer_class', None):
            self.assertEqual(self.client.get('/api/system/user/export_data/').status_code, 405)

    def test_import_and_log_workbooks_remain_shutdown_before_file_read(self):
        self.client.force_authenticate(self.admin)
        with patch('coreadmin.utils.import_export_mixin.import_to_data') as importer, patch(
                'coreadmin.utils.import_export.openpyxl.load_workbook') as reader:
            for resource in ('user', 'operation_log', 'login_log'):
                for method in ('get', 'post', 'head'):
                    self.assertEqual(getattr(self.client, method)(f'/api/system/{resource}/import_data/').status_code, 405)
            for resource in ('operation_log', 'login_log'):
                for action in ('export_data', 'update_template'):
                    self.assertEqual(self.client.get(f'/api/system/{resource}/{action}/').status_code, 405)
            for method in ('post', 'head'):
                self.assertEqual(getattr(self.client, method)('/api/system/dept/import_data/', {'url': 'never-open.xlsx'}).status_code, 405)
            importer.assert_not_called()
            reader.assert_not_called()
