from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from coreadmin.foundation_tests.access_fixtures import grant
from coreadmin.system.models import (Users, Dept, Role, Menu, MenuButton, ApiWhiteList,
    FieldPermission, MessageCenter, MessageCenterTargetUser)


class ObjectFieldPolicyTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.dept = Dept.objects.create(name='One')
        cls.other_dept = Dept.objects.create(name='Two', parent=cls.dept)
        cls.actor = Users.objects.create(username='actor', name='Actor', dept=cls.dept, pwd_change_count=1)
        cls.admin = Users.objects.create(username='admin', is_superuser=True, pwd_change_count=1)
        cls.a = Users.objects.create(username='a', name='Visible A', email='a@example.test', creator=cls.actor, dept_belong_id=cls.dept.id, dept=cls.dept, pwd_change_count=1)
        cls.b = Users.objects.create(username='b', name='Hidden B', email='b@example.test', dept_belong_id=cls.other_dept.id, dept=cls.other_dept, pwd_change_count=1)
        cls.client_user = APIClient()
        cls.client_user.force_authenticate(cls.actor)

    def test_unknown_alias_and_whitelist_cannot_grant_action(self):
        grant(self.actor, 'unknown', 'Users', fields=['name'])
        ApiWhiteList.objects.create(url='/api/system/user/', method=0, enable_datasource=False)
        self.assertEqual(self.client_user.get('/api/system/user/').status_code, 403)

    def test_valid_alias_ignores_incorrect_url_method_and_has_positive_read(self):
        grant(self.actor, 'user:Retrieve', 'Users', fields=['name'], scope=0)
        response = self.client_user.get(f'/api/system/user/{self.a.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['data'], {'id': self.a.id, 'name': self.a.name})
        self.assertEqual(self.client_user.get(f'/api/system/user/{self.b.id}/').status_code, 404)

    def test_disabled_role_next_request_denies(self):
        permission = grant(self.actor, 'user:Retrieve', 'Users', fields=['name'])
        self.assertEqual(self.client_user.get(f'/api/system/user/{self.a.id}/').status_code, 200)
        Role.objects.filter(pk=permission.role_id).update(status=False)
        self.assertEqual(self.client_user.get(f'/api/system/user/{self.a.id}/').status_code, 403)

    def test_show_all_preserves_scope_and_protected_user_envelope(self):
        grant(self.actor, 'user:Search', 'Users', fields=['name', 'dept'], scope=0)
        self.admin.dept = self.other_dept
        self.admin.creator = self.actor
        self.admin.save()
        response = self.client_user.get('/api/system/user/', {'dept': self.dept.id, 'show_all': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual({row['id'] for row in response.data['data']}, {self.a.id})

    def test_per_object_field_tuples_do_not_cross_scopes(self):
        grant(self.actor, 'user:Search', 'Users', fields=['name', 'email'], scope=2)
        grant(self.actor, 'user:Search', 'Users', fields=['name'], scope=4, departments=[self.other_dept])
        response = self.client_user.get('/api/system/user/')
        self.assertEqual(response.status_code, 200)
        rows = {row['id']: row for row in response.data['data']}
        self.assertEqual(rows[self.a.id]['email'], self.a.email)
        self.assertNotIn('email', rows[self.b.id])

    def test_no_field_grant_returns_only_explicit_identity(self):
        grant(self.actor, 'user:Retrieve', 'Users')
        response = self.client_user.get(f'/api/system/user/{self.a.id}/')
        self.assertEqual(response.data['data'], {'id': self.a.id})

    def test_hidden_filter_search_order_restql_rejected(self):
        grant(self.actor, 'user:Search', 'Users', fields=['name'])
        for query in [{'email': self.a.email}, {'ordering': 'email'}, {'search': 'a'},
                      {'query': '{id,email}'}, {'dept__name': 'One'}, {'query': '{role_info{id}}'}]:
            with self.subTest(query=query):
                self.assertEqual(self.client_user.get('/api/system/user/', query).status_code, 400)
        self.assertEqual(self.client_user.get('/api/system/user/', {'name': 'Visible', 'ordering': 'name', 'query': '{id,name}'}).status_code, 200)

    def test_update_patch_and_head(self):
        grant(self.actor, 'user:Update', 'Users', fields=['name'], update=['username', 'name'], scope=0)
        for method in ['put', 'patch']:
            with self.subTest(method=method):
                self.assertEqual(getattr(self.client_user, method)(f'/api/system/user/{self.a.id}/', {'username': self.a.username, 'name': method}, format='json').status_code, 200)
                self.assertEqual(getattr(self.client_user, method)(f'/api/system/user/{self.b.id}/', {'name': method}, format='json').status_code, 404)
                self.assertEqual(getattr(self.client_user, method)(f'/api/system/user/{self.a.id}/', {'email': 'no@example.test'}, format='json').status_code, 400)
        grant(self.actor, 'user:Retrieve', 'Users', fields=['name'], scope=0)
        self.assertEqual(self.client_user.head(f'/api/system/user/{self.a.id}/').status_code, 200)
        self.assertEqual(self.client_user.options('/api/system/user/').data['data'], {})

    def test_create_scope_and_field_ceiling_bulk_zero_partial_write(self):
        grant(self.actor, 'user:Create', 'Users', fields=['name'], create=['username', 'name', 'password', 'is_superuser'], scope=0)
        response = self.client_user.post('/api/system/user/', {'username': 'new', 'name': 'New', 'password': 'Valid-test-password-2026!'}, format='json')
        self.assertEqual(response.status_code, 200)
        created = Users.objects.get(username='new')
        self.assertEqual(created.creator_id, self.actor.id)
        count = Users.objects.count()
        response = self.client_user.post('/api/system/user/', [
            {'username': 'bulk-safe'}, {'username': 'bulk-unsafe', 'is_superuser': True}], format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Users.objects.count(), count)

    def test_create_scope_without_matching_target_is_denied(self):
        grant(self.actor, 'user:Create', 'Users', create=['username'], scope=4, departments=[self.other_dept])
        self.assertEqual(self.client_user.post('/api/system/user/', {'username': 'out'}, format='json').status_code, 403)
        self.assertFalse(Users.objects.filter(username='out').exists())

    def test_bulk_missing_target_does_not_partially_delete(self):
        client = APIClient(); client.force_authenticate(self.admin)
        role = Role.objects.create(name='Delete', key='delete')
        response = client.delete('/api/system/role/multiple_delete/', {'keys': [role.id, 999999]}, format='json')
        self.assertEqual(response.status_code, 404)
        self.assertTrue(Role.objects.filter(pk=role.id).exists())
        self.assertEqual(client.delete('/api/system/role/multiple_delete/', {'keys': [role.id, role.id]}, format='json').status_code, 200)

    def test_custom_admin_targets_all_validated_before_write(self):
        client = APIClient(); client.force_authenticate(self.admin)
        role = Role.objects.create(name='Target', key='target')
        response = client.post(f'/api/system/role/{role.id}/add_role_users/', {'users_id': [self.actor.id, 999999]}, format='json')
        self.assertEqual(response.status_code, 404)
        self.assertFalse(role.users_set.exists())

    def test_internal_object_error_is_500_not_404(self):
        client = APIClient(); client.force_authenticate(self.admin)
        with patch('coreadmin.system.views.user.UserViewSet.filter_queryset', side_effect=RuntimeError('B2_INTERNAL')):
            response = client.get(f'/api/system/user/{self.a.id}/')
        self.assertEqual(response.status_code, 500)
        self.assertNotIn('B2_INTERNAL', str(response.data))

    def test_message_recipient_and_creator_detail_and_read_mark(self):
        payload = '<script>B2</script>\n<b>plain text</b>'
        own = MessageCenter.objects.create(title='Own', content=payload, creator=self.admin)
        other = MessageCenter.objects.create(title='Other', content='other', creator=self.admin)
        sent = MessageCenter.objects.create(title='Sent', content='sent', creator=self.actor)
        relation = MessageCenterTargetUser.objects.create(messagecenter=own, users=self.actor, is_read=False)
        grant(self.actor, 'messageCenter:Retrieve', 'MessageCenter', fields=['title', 'content'], scope=3)
        self.assertEqual(self.client_user.get(f'/api/system/message_center/{other.id}/').status_code, 404)
        relation.refresh_from_db(); self.assertFalse(relation.is_read)
        response = self.client_user.get(f'/api/system/message_center/{own.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['data']['content'], payload)
        relation.refresh_from_db(); self.assertTrue(relation.is_read)
        self.assertEqual(self.client_user.get(f'/api/system/message_center/{sent.id}/').status_code, 200)
        for action in ['get_self_receive', 'get_newest_msg']:
            self.assertEqual(self.client_user.get(f'/api/system/message_center/{action}/').status_code, 200)
        self.assertEqual(self.client_user.post('/api/system/message_center/', {'title': 'No'}, format='json').status_code, 403)

    def test_export_both_routes_share_scoped_field_output(self):
        grant(self.actor, 'user:Export', 'Users', fields=['username'], scope=0)
        for method, url in [('get', '/api/system/user/export_data/'), ('post', '/api/system/user/export/')]:
            with self.subTest(method=method):
                response = getattr(self.client_user, method)(url)
                self.assertEqual(response.status_code, 200)
                rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[1], (1, self.a.username))
                self.assertNotIn(self.a.email, str(rows))
                self.assertNotIn(self.b.username, str(rows))

    def test_template_data_is_not_wider_than_field_policy(self):
        grant(self.actor, 'user:UpdateTemplate', 'Users', fields=['username'], scope=0)
        response = self.client_user.get('/api/system/user/update_template/')
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertEqual(rows[1], (1, self.a.id, self.a.username))
        self.assertNotIn(self.a.email, str(rows))

    def test_nested_role_cannot_borrow_parent_user_field_grant(self):
        target_role = Role.objects.create(name='Nested role', key='nested-key')
        self.a.role.add(target_role)
        grant(self.actor, 'user:Retrieve', 'Users', fields=['name', 'role_info'])
        url = f'/api/system/user/{self.a.id}/'
        response = self.client_user.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['data']['role_info'], [])
        grant(self.actor, 'role:Retrieve', 'Role', fields=['name'])
        response = self.client_user.get(url)
        self.assertEqual(response.data['data']['role_info'], [{'id': target_role.id, 'name': target_role.name}])
        self.assertNotIn(target_role.key, str(response.data))

    def test_update_field_from_other_scope_cannot_be_borrowed(self):
        grant(self.actor, 'user:Update', 'Users', fields=['name'], update=['email'], scope=2)
        grant(self.actor, 'user:Update', 'Users', fields=['name'], update=['name'], scope=4, departments=[self.other_dept])
        response = self.client_user.patch(f'/api/system/user/{self.b.id}/', {'email': 'changed@example.test'}, format='json')
        self.assertEqual(response.status_code, 400)
        self.b.refresh_from_db(); self.assertEqual(self.b.email, 'b@example.test')
        self.assertEqual(self.client_user.patch(f'/api/system/user/{self.b.id}/', {'name': 'Allowed'}, format='json').status_code, 200)

    def test_other_action_field_role_cannot_contribute(self):
        grant(self.actor, 'user:Retrieve', 'Users', fields=['name'])
        grant(self.actor, 'user:Search', 'Users', fields=['email'])
        response = self.client_user.get(f'/api/system/user/{self.a.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('email', response.data['data'])

    def test_delete_object_scope_positive_and_negative(self):
        grant(self.actor, 'user:Delete', 'Users', scope=0)
        self.assertEqual(self.client_user.delete(f'/api/system/user/{self.b.id}/').status_code, 404)
        self.assertTrue(Users.objects.filter(pk=self.b.id).exists())
        self.assertEqual(self.client_user.delete(f'/api/system/user/{self.a.id}/').status_code, 200)
        self.assertFalse(Users.objects.filter(pk=self.a.id).exists())

    def test_real_jwt_disabled_role_revocation(self):
        from coreadmin.system.services.auth import AuthService
        permission = grant(self.actor, 'user:Retrieve', 'Users', fields=['name'])
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION='JWT ' + AuthService.issue(self.actor)['access'])
        url = f'/api/system/user/{self.a.id}/'
        self.assertEqual(client.get(url).status_code, 200)
        Role.objects.filter(pk=permission.role_id).update(status=False)
        self.assertEqual(client.get(url).status_code, 403)

    def test_empty_create_cannot_bypass_default_deny_fields(self):
        grant(self.actor, 'menu:Create', 'Menu', scope=3)
        before = Menu.objects.count()
        self.assertEqual(self.client_user.post('/api/system/menu/', {}, format='json').status_code, 403)
        self.assertEqual(Menu.objects.count(), before)

    def test_readonly_unknown_query_never_mutates_serializer_class(self):
        grant(self.actor, 'user:Retrieve', 'Users', fields=['name'])
        normal = self.client_user.get(f'/api/system/user/{self.a.id}/')
        self.assertNotIn('email', normal.data['data'])
        client = APIClient(); client.force_authenticate(self.admin)
        admin = client.get(f'/api/system/user/{self.a.id}/')
        self.assertEqual(admin.data['data']['email'], self.a.email)
        self.assertEqual(client.get('/api/system/user/', {'password': 'not-readable'}).status_code, 400)

    def test_create_target_matches_saved_attribution_even_with_restql(self):
        grant(self.actor, 'menu:Create', 'Menu', fields=['name'], create=['name'], scope=0)
        response = self.client_user.post('/api/system/menu/?query={name}', {'name': 'Scoped creation'}, format='json')
        self.assertEqual(response.status_code, 200)
        obj = Menu.objects.get(name='Scoped creation')
        self.assertEqual(obj.creator_id, self.actor.pk)
        self.assertEqual(str(obj.dept_belong_id), str(self.actor.dept_id))
        grant(self.actor, 'user:Create', 'Users', create=['username'], scope=2)
        # B1 generic User create cannot assign a department. A DEPT grant
        # therefore cannot claim this proposed target, despite the actor's dept.
        self.assertEqual(self.client_user.post('/api/system/user/', {'username': 'dept-only'}, format='json').status_code, 403)
        self.assertFalse(Users.objects.filter(username='dept-only').exists())

    def test_approved_search_positive_paths(self):
        grant(self.actor, 'role:Search', 'Role', fields=['name', 'key'])
        self.assertEqual(self.client_user.get('/api/system/role/', {'search': 'Access'}).status_code, 200)
        client = APIClient(); client.force_authenticate(self.admin)
        self.assertEqual(client.get('/api/system/user/', {'search': self.a.name}).status_code, 200)
